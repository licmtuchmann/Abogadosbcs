#!/usr/bin/env python3
"""Genera el formato de ENTREVISTA A TESTIGO del despacho
Meza Figueroa, Tuchmann & Asociados.

Uso:
    # Formato modelo en blanco (para llenar a mano)
    python3 generar_entrevista.py --salida ENTREVISTA_TESTIGO_PLANTILLA.docx

    # Formato llenado total o parcialmente con datos de un archivo JSON
    python3 generar_entrevista.py --datos datos.json --salida entrevista.docx

Todo campo ausente o vacío en el JSON se imprime como línea en blanco,
de modo que el mismo documento puede completarse después a mano.
"""
import argparse
import json
import os
from datetime import date

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

NAVY = RGBColor(0x0B, 0x3D, 0x6E)
FUENTE = "Arial"
DESPACHO = "MEZA FIGUEROA, TUCHMANN & ASOCIADOS"
AQUI = os.path.dirname(os.path.abspath(__file__))
LOGO = os.path.join(AQUI, "..", "assets", "logo-mark.png")

ANCHO_UTIL = Cm(21.59 - 2 * 2.0)  # carta, márgenes de 2 cm

CIERRE = (
    "Leída que fue la presente entrevista por la persona entrevistada, "
    "manifiesta que lo asentado corresponde fielmente a lo expresado por ella "
    "de manera libre y voluntaria, y la firma al margen y al calce para constancia."
)

# ---------------------------------------------------------------- números
UNIDADES = ["cero", "uno", "dos", "tres", "cuatro", "cinco", "seis", "siete",
            "ocho", "nueve", "diez", "once", "doce", "trece", "catorce",
            "quince", "dieciséis", "diecisiete", "dieciocho", "diecinueve",
            "veinte", "veintiuno", "veintidós", "veintitrés", "veinticuatro",
            "veinticinco", "veintiséis", "veintisiete", "veintiocho",
            "veintinueve"]
DECENAS = ["", "", "", "treinta", "cuarenta", "cincuenta", "sesenta",
           "setenta", "ochenta", "noventa"]
CENTENAS = ["", "ciento", "doscientos", "trescientos", "cuatrocientos",
            "quinientos", "seiscientos", "setecientos", "ochocientos",
            "novecientos"]
MESES = ["enero", "febrero", "marzo", "abril", "mayo", "junio", "julio",
         "agosto", "septiembre", "octubre", "noviembre", "diciembre"]


def en_letra(n):
    """Número entero de 0 a 9999 en letra (español de México)."""
    if n < 30:
        return UNIDADES[n]
    if n < 100:
        d, u = divmod(n, 10)
        return DECENAS[d] + ("" if u == 0 else " y " + UNIDADES[u])
    if n < 1000:
        if n == 100:
            return "cien"
        c, r = divmod(n, 100)
        return CENTENAS[c] + ("" if r == 0 else " " + en_letra(r))
    m, r = divmod(n, 1000)
    miles = "mil" if m == 1 else en_letra(m) + " mil"
    return miles + ("" if r == 0 else " " + en_letra(r))


def fecha_forense(iso):
    """'2026-06-12' -> '12 (doce) de junio de 2026 (dos mil veintiséis)'."""
    f = date.fromisoformat(iso)
    dia = "primero" if f.day == 1 else en_letra(f.day)
    return (f"{f.day} ({dia}) de {MESES[f.month - 1]} de "
            f"{f.year} ({en_letra(f.year)})")


def hora_forense(hhmm):
    """'11:40' -> '11:40 (once horas con cuarenta minutos)'."""
    h, m = (int(x) for x in hhmm.split(":"))
    texto = "una hora" if h == 1 else f"{en_letra(h)} horas"
    if m:
        texto += " con " + ("un minuto" if m == 1 else f"{en_letra(m)} minutos")
    return f"{h:02d}:{m:02d} ({texto})"


def fecha_y_hora(datos):
    partes = []
    if datos.get("hora"):
        partes.append(hora_forense(datos["hora"]))
    if datos.get("fecha"):
        partes.append(("del " if partes else "") + fecha_forense(datos["fecha"]))
    return " ".join(partes)


def edad_forense(valor):
    if valor in (None, ""):
        return ""
    try:
        n = int(valor)
    except (TypeError, ValueError):
        return str(valor)
    return f"{n} ({en_letra(n)}) años"


# ---------------------------------------------------------------- formato
def run(p, texto, negrita=False, tam=12, color=None):
    r = p.add_run(texto)
    r.font.name = FUENTE
    r._element.rPr.rFonts.set(qn("w:eastAsia"), FUENTE)
    r.font.size = Pt(tam)
    r.bold = negrita
    if color:
        r.font.color.rgb = color
    return r


def parrafo(contenedor, alinear=WD_ALIGN_PARAGRAPH.JUSTIFY, antes=0, despues=4):
    p = contenedor.add_paragraph()
    p.alignment = alinear
    pf = p.paragraph_format
    pf.space_before = Pt(antes)
    pf.space_after = Pt(despues)
    pf.line_spacing = 1.15
    return p


def campo(contenedor, etiqueta, valor, ancho=ANCHO_UTIL):
    """ETIQUETA: valor  — o ETIQUETA: ________ hasta el margen si está vacío."""
    p = parrafo(contenedor, WD_ALIGN_PARAGRAPH.LEFT, despues=6)
    run(p, etiqueta + ": ", negrita=True, color=NAVY)
    if valor:
        run(p, str(valor))
    else:
        # Tabulador izquierdo con guía de subrayado: traza la línea en blanco
        # hasta el borde del renglón o de la celda.
        p.paragraph_format.tab_stops.add_tab_stop(
            ancho, WD_TAB_ALIGNMENT.LEFT, WD_TAB_LEADER.LINES)
        run(p, "\t")
    return p


def lineas_en_blanco(contenedor, n):
    for _ in range(n):
        p = parrafo(contenedor, WD_ALIGN_PARAGRAPH.LEFT, antes=6, despues=0)
        p.paragraph_format.tab_stops.add_tab_stop(
            ANCHO_UTIL, WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.LINES)
        run(p, "\t")


def sin_bordes(tabla):
    tblPr = tabla._tbl.tblPr
    bordes = OxmlElement("w:tblBorders")
    for lado in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement(f"w:{lado}")
        e.set(qn("w:val"), "nil")
        bordes.append(e)
    tblPr.append(bordes)


def fila_de_campos(doc, pares, pesos=None):
    """Varios campos en un mismo renglón (p. ej. correo / edad).

    Se arma con tabuladores (no con tabla) para que las líneas en blanco se
    vean igual en Word y en LibreOffice. Si algún valor es largo, cada campo
    pasa a su propio renglón para evitar encimados.
    """
    if any(len(str(v or "")) > 22 for _, v in pares):
        for etiqueta, valor in pares:
            campo(doc, etiqueta, valor)
        return
    pesos = pesos or [1] * len(pares)
    p = parrafo(doc, WD_ALIGN_PARAGRAPH.LEFT, despues=6)
    acumulado = 0
    for i, ((etiqueta, valor), w) in enumerate(zip(pares, pesos)):
        acumulado += w
        fin = int(ANCHO_UTIL * acumulado / sum(pesos))
        ultimo = i == len(pares) - 1
        tope = fin if ultimo else fin - Cm(0.4)
        run(p, etiqueta + ": ", negrita=True, color=NAVY)
        if valor:
            run(p, str(valor))
            if not ultimo:
                p.paragraph_format.tab_stops.add_tab_stop(fin, WD_TAB_ALIGNMENT.LEFT)
                run(p, "\t")
        else:
            p.paragraph_format.tab_stops.add_tab_stop(
                tope, WD_TAB_ALIGNMENT.LEFT, WD_TAB_LEADER.LINES)
            run(p, "\t")
            if not ultimo:
                p.paragraph_format.tab_stops.add_tab_stop(fin, WD_TAB_ALIGNMENT.LEFT)
                run(p, "\t")


def seccion(doc, titulo):
    p = parrafo(doc, WD_ALIGN_PARAGRAPH.LEFT, antes=10, despues=4)
    run(p, titulo, negrita=True, color=NAVY)
    borde = OxmlElement("w:pBdr")
    inf = OxmlElement("w:bottom")
    for k, v in (("w:val", "single"), ("w:sz", "6"), ("w:space", "1"),
                 ("w:color", "0B3D6E")):
        inf.set(qn(k), v)
    borde.append(inf)
    p._p.get_or_add_pPr().append(borde)


def campo_numero_pagina(p, instruccion):
    for tipo, texto in (("begin", None), (None, instruccion), ("end", None)):
        r = run(p, "", tam=9, color=NAVY)
        if tipo:
            e = OxmlElement("w:fldChar")
            e.set(qn("w:fldCharType"), tipo)
        else:
            e = OxmlElement("w:instrText")
            e.set(qn("xml:space"), "preserve")
            e.text = texto
        r._r.append(e)


def encabezado_y_pie(doc):
    sec = doc.sections[0]
    enc = sec.header
    tabla = enc.add_table(rows=1, cols=2, width=ANCHO_UTIL)
    sin_bordes(tabla)
    c_logo, c_txt = tabla.rows[0].cells
    c_logo.width, c_txt.width = Cm(2.0), ANCHO_UTIL - Cm(2.0)
    if os.path.exists(LOGO):
        c_logo.paragraphs[0].add_run().add_picture(LOGO, height=Cm(1.5))
    p = c_txt.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run(p, DESPACHO, negrita=True, tam=10, color=NAVY)
    p2 = c_txt.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run(p2, "Documento confidencial — Entrevista a testigo", tam=8, color=NAVY)
    enc.paragraphs[0]._element.getparent().remove(enc.paragraphs[0]._element)

    pie = sec.footer.paragraphs[0]
    pie.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run(pie, "Página ", tam=9, color=NAVY)
    campo_numero_pagina(pie, "PAGE")
    run(pie, " de ", tam=9, color=NAVY)
    campo_numero_pagina(pie, "NUMPAGES")


def firmas(doc, entrevistado, entrevistador):
    parrafo(doc, despues=24).paragraph_format.keep_with_next = True
    tabla = doc.add_table(rows=3, cols=2)
    sin_bordes(tabla)
    tabla.alignment = WD_TABLE_ALIGNMENT.CENTER
    filas = (("ENTREVISTADO(A)", "ENTREVISTADOR(A)"),
             ("_______________________________", "_______________________________"),
             (entrevistado or "Nombre y firma", entrevistador or "Nombre y firma"))
    for i, (izq, der) in enumerate(filas):
        for celda, texto in zip(tabla.rows[i].cells, (izq, der)):
            p = celda.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(36 if i == 1 else 2)
            p.paragraph_format.keep_with_next = i < 2  # firmas en una sola hoja
            run(p, texto, negrita=(i != 1), tam=10 if i == 2 else 12,
                color=NAVY if i == 0 else None)
        trPr = tabla.rows[i]._tr.get_or_add_trPr()
        trPr.append(OxmlElement("w:cantSplit"))


# ---------------------------------------------------------------- documento
def generar(datos, salida, preguntas_en_blanco=10, lineas_relato=12):
    d = {k: (v.strip() if isinstance(v, str) else v) for k, v in datos.items()}
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.59), Cm(27.94)
    for lado in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(sec, lado, Cm(2.0))
    estilo = doc.styles["Normal"]
    estilo.font.name = FUENTE
    estilo.font.size = Pt(12)
    estilo.element.rPr.rFonts.set(qn("w:eastAsia"), FUENTE)
    encabezado_y_pie(doc)

    p = parrafo(doc, WD_ALIGN_PARAGRAPH.CENTER, despues=2)
    run(p, "ENTREVISTA A TESTIGO", negrita=True, tam=14, color=NAVY)
    campo(doc, d.get("etiqueta_identificador") or "NÚMERO ÚNICO DE CAUSA (NUC)",
          d.get("nuc"))
    if d.get("asunto"):
        campo(doc, "ASUNTO", d["asunto"])

    seccion(doc, "I. DATOS DE LA DILIGENCIA")
    campo(doc, "LUGAR", d.get("lugar"))
    campo(doc, "FECHA Y HORA DE INICIO", fecha_y_hora(d))
    campo(doc, "ENTREVISTADOR(A)", d.get("entrevistador"))
    campo(doc, "ENTREVISTADO(A)", d.get("entrevistado"))

    seccion(doc, "II. DATOS GENERALES DE LA PERSONA ENTREVISTADA")
    campo(doc, "DOMICILIO DEL TESTIGO", d.get("domicilio"))
    campo(doc, "TIEMPO DE RESIDENCIA EN ESTA CIUDAD", d.get("tiempo_residencia"))
    fila_de_campos(doc, [("CORREO ELECTRÓNICO", d.get("correo")),
                         ("EDAD", edad_forense(d.get("edad")))], [3, 2])
    campo(doc, "MEDIA FILIACIÓN", d.get("media_filiacion"))
    fila_de_campos(doc, [("APODO", d.get("apodo")),
                         ("SEXO", d.get("sexo")),
                         ("TELÉFONO CELULAR", d.get("telefono"))], [3, 3, 4])
    campo(doc, "OCUPACIÓN", d.get("ocupacion"))
    campo(doc, "DOMICILIO LABORAL", d.get("domicilio_laboral"))
    fila_de_campos(doc, [("IDENTIFICACIÓN (TIPO)", d.get("identificacion_tipo")),
                         ("NÚMERO", d.get("identificacion_numero"))], [3, 2])
    campo(doc, "RELACIÓN CON LAS PARTES O LOS HECHOS", d.get("relacion"))

    seccion(doc, "III. RELATO DEL HECHO")
    relato = d.get("relato")
    if relato:
        for bloque in [b for b in relato.split("\n") if b.strip()]:
            p = parrafo(doc, despues=8)
            run(p, bloque.strip())
    else:
        lineas_en_blanco(doc, lineas_relato)

    seccion(doc, "IV. PREGUNTAS Y RESPUESTAS")
    preguntas = d.get("preguntas") or [{} for _ in range(preguntas_en_blanco)]
    for i, q in enumerate(preguntas, 1):
        if isinstance(q, str):
            q = {"pregunta": q}
        p = parrafo(doc, WD_ALIGN_PARAGRAPH.JUSTIFY, antes=6, despues=2)
        run(p, f"PREGUNTA {i}: ", negrita=True, color=NAVY)
        if q.get("pregunta"):
            run(p, q["pregunta"])
        else:
            p.paragraph_format.tab_stops.add_tab_stop(
                ANCHO_UTIL, WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.LINES)
            run(p, "\t")
        p = parrafo(doc, WD_ALIGN_PARAGRAPH.JUSTIFY, despues=2)
        run(p, "RESPUESTA: ", negrita=True, color=NAVY)
        if q.get("respuesta"):
            run(p, q["respuesta"])
        else:
            p.paragraph_format.tab_stops.add_tab_stop(
                ANCHO_UTIL, WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.LINES)
            run(p, "\t")
            lineas_en_blanco(doc, int(q.get("lineas_respuesta", 2)))

    seccion(doc, "V. CIERRE")
    campo(doc, "HORA DE CONCLUSIÓN",
          hora_forense(d["hora_conclusion"]) if d.get("hora_conclusion") else "")
    p = parrafo(doc, despues=4, antes=6)
    p.paragraph_format.keep_with_next = True
    run(p, d.get("cierre") or CIERRE)

    firmas(doc, d.get("entrevistado"), d.get("entrevistador"))
    doc.core_properties.title = "Entrevista a testigo"
    doc.core_properties.author = DESPACHO
    doc.save(salida)
    return salida


def main():
    ap = argparse.ArgumentParser(description=__doc__,
                                 formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--datos", help="archivo JSON con los datos (opcional)")
    ap.add_argument("--salida", required=True, help="ruta del .docx a generar")
    ap.add_argument("--preguntas-en-blanco", type=int, default=10,
                    help="espacios de pregunta si el JSON no trae preguntas")
    ap.add_argument("--lineas-relato", type=int, default=12,
                    help="renglones en blanco para el relato si viene vacío")
    a = ap.parse_args()
    datos = {}
    if a.datos:
        with open(a.datos, encoding="utf-8") as f:
            datos = json.load(f)
    print(generar(datos, a.salida, a.preguntas_en_blanco, a.lineas_relato))


if __name__ == "__main__":
    main()
